import openpyxl
import sys

def print_excel_structure(file_path):
    print(f"--- Structure of {file_path} ---")
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        print("Sheet names:", wb.sheetnames)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            print(f"\nSheet: {sheet_name}")
            for row in sheet.iter_rows(min_row=1, max_row=5, values_only=True):
                print(row)
    except Exception as e:
        print(f"Error reading {file_path}: {e}")

if __name__ == "__main__":
    print_excel_structure("Bitacoras de Mercado local abril.xlsx")
    print_excel_structure("DOC-20260112-WA0008..xlsx")
